from io import BytesIO
from contextlib import closing
import json
from unittest.mock import Mock

import openpyxl
from PIL import Image
import pypdfium2 as pdfium
import pytest
from reportlab.pdfgen import canvas
from werkzeug.datastructures import FileStorage
from werkzeug.test import Client
from werkzeug.wrappers import Response

from src import config, estoque, cerebro, documentos_estoque as docs, catalogo_pdf
from src.db import init_db, get_conn


@pytest.fixture
def loja(tmp_path, monkeypatch):
    monkeypatch.setattr(config,'DATABASE_PATH',str(tmp_path/'loja.db'))
    monkeypatch.setattr(config,'SEED_CATALOGO',False)
    monkeypatch.setattr(config,'MODO_TESTE',True)
    monkeypatch.setattr(config,'OLLAMA_URL','http://localhost:11434')
    init_db()
    from src.app import app
    monkeypatch.setitem(app.config,'DESKTOP_TOKEN','documento-local')
    c=app.test_client()
    c.get('/gestao?acesso=documento-local',follow_redirects=True)
    with c.session_transaction() as s:
        csrf=s['csrf']
    return app,c,{'X-CSRF-Token':csrf}


def arquivo(texto='nome;categoria;preco;sku;unidade;estoque\nCaderno;Papelaria;15,90;CAD;un;5\n',nome='estoque.csv'):
    return FileStorage(stream=BytesIO(texto.encode() if isinstance(texto,str) else texto),filename=nome)


def lancar(revisao, **campos):
    return docs.aplicar({'token':revisao['token'],'produtos':revisao['produtos'],**campos},'dono')


def texto_pdf(dados):
    with pdfium.PdfDocument(dados) as pdf:
        textos=[]
        for i in range(len(pdf)):
            with closing(pdf[i]) as p, closing(p.get_textpage()) as t:
                textos.append(t.get_text_range())
        return '\n'.join(textos)


def test_csv_previa_sem_mudar_estoque_lancamento_e_obsidian(loja):
    r=docs.ler(arquivo(),'dono')
    assert not r['usou_ollama'] and estoque.listar_produtos()==[]
    result=lancar(r)
    assert result['nota_salva']
    p=estoque.listar_produtos()[0]
    assert p['estoque']==5 and p['preco']==15.90
    notas=list((cerebro.diretorio()/'Documentos de estoque').glob('*.md'))
    assert len(notas)==1 and 'Caderno' in notas[0].read_text(encoding='utf-8')
    assert 'Caderno' not in cerebro.instrucoes_contexto('caderno')
    assert 'Documentos de estoque/'+notas[0].name in cerebro.snapshot()
    assert lancar(r)['ja_aplicado'] and estoque.listar_produtos()[0]['estoque']==5
    with pytest.raises(ValueError,match='já foi lançado'):
        docs.ler(arquivo(),'dono')


def test_xlsx_compra_custo_nao_vira_preco_venda(loja):
    book=openpyxl.Workbook()
    book.active.append(['Produto','Categoria','Valor unitário','Quantidade'])
    book.active.append(['Café','Mercearia',12.5,3])
    out=BytesIO();book.save(out);book.close()
    r=docs.ler(arquivo(out.getvalue(),'nota.xlsx'),'dono',tipo='compra')
    assert r['produtos'][0]['custo']=='12.5' and r['produtos'][0]['preco']==''
    with pytest.raises(ValueError,match='Preço'):
        lancar(r)
    assert estoque.listar_produtos()==[]
    r['produtos'][0]['preco']='19,90'
    lancar(r,obsidian=False)
    assert estoque.listar_produtos()[0]['custo']==12.5


def test_entrada_preserva_preco_foto_e_contagem_detecta_venda(loja):
    p=estoque.salvar_produto({'nome':'Caderno','categoria':'Papelaria','preco':15.9,'sku':'CAD','estoque':8})
    r=docs.ler(arquivo(),'dono',tipo='compra')
    assert r['produtos'][0]['acao']=='entrada'
    lancar(r)
    assert estoque.listar_produtos()[0]['estoque']==13
    r=docs.ler(arquivo(),'dono',tipo='lista')
    assert r['produtos'][0]['acao']=='contagem'
    estoque.ajustar_estoque(p['id'],-1,'Venda depois da leitura')
    with pytest.raises(ValueError,match='saldo mudou'):
        lancar(r)
    assert estoque.listar_produtos()[0]['estoque']==12
    r=docs.ler(arquivo(),'dono')
    lancar(r)
    assert estoque.listar_produtos()[0]['estoque']==5 and estoque.listar_produtos()[0]['preco']==15.9


def test_falha_de_linha_desfaz_tudo_e_marca_so_apos_confirmacao(loja):
    r=docs.ler(arquivo('nome;categoria;preco;sku;estoque\nA;Itens;10;A;2\nB;Itens;-1;B;3'),'dono')
    with pytest.raises(ValueError):
        lancar(r)
    assert estoque.listar_produtos()==[]
    with get_conn() as c:
        assert c.execute('SELECT concluido FROM documentos_estoque').fetchone()[0]==0
    r['produtos'][1]['preco']='20'
    assert lancar(r)['quantidade']==2


def test_outros_donos_expiracao_duplicidade_e_unidades(loja):
    r=docs.ler(arquivo(),'dono')
    with pytest.raises(ValueError,match='outra sessão'):
        docs.aplicar({'token':r['token'],'produtos':r['produtos']},'outra')
    with get_conn() as c:
        c.execute("UPDATE documentos_estoque SET criado_em=datetime('now','-3 hours')")
    with pytest.raises(ValueError,match='expirou'):
        lancar(r)
    p=estoque.salvar_produto({'nome':'Caderno','categoria':'Itens','preco':10,'sku':'CAD','estoque':3})
    r=docs.ler(arquivo(),'dono',tipo='compra');r['produtos'][0]['unidade']='kg'
    with pytest.raises(ValueError,match='unidade'):
        lancar(r)
    assert estoque.listar_produtos()[0]['estoque']==3


def test_duplicado_entre_revisoes_nao_soma_duas_vezes(loja):
    estoque.salvar_produto({'nome':'Caderno','categoria':'Itens','preco':10,'sku':'CAD','estoque':3})
    a=docs.ler(arquivo(),'dono',tipo='compra');b=docs.ler(arquivo(),'dono',tipo='compra')
    lancar(a)
    with pytest.raises(ValueError,match='outra revisão'):
        lancar(b)
    assert estoque.listar_produtos()[0]['estoque']==8


def test_pdf_texto_e_nota_md_enviados_ao_modelo_com_revisao(loja,monkeypatch):
    out=BytesIO();c=canvas.Canvas(out);c.drawString(40,700,'Caderno 5 unidades Preco de venda R$ 15,90');c.save()
    fake=Mock(return_value=[{'nome':'Caderno','categoria':'Papelaria','preco':15.9,'estoque':5}])
    monkeypatch.setattr(docs,'_extrair_ollama',fake)
    r=docs.ler(arquivo(out.getvalue(),'produtos.pdf'),'dono')
    assert r['usou_ollama'] and 'Caderno' in fake.call_args.args[0]
    docs.ler(arquivo('# Lista\nCaderno: 5 unidades', 'nota.md'),'dono')
    assert fake.call_count==2 and not estoque.listar_produtos()


def test_pdf_digitalizado_produz_imagens_para_modelo(loja,monkeypatch):
    out=BytesIO();c=canvas.Canvas(out);c.rect(10,10,100,200,fill=1);c.save()
    fake=Mock(return_value=[{'nome':'Papel','estoque':2}]);monkeypatch.setattr(docs,'_extrair_ollama',fake)
    docs.ler(arquivo(out.getvalue(),'scan.pdf'),'dono')
    assert len(fake.call_args.args[1])==1


def test_modelo_local_sem_ferramentas_e_resposta_limitada(loja,monkeypatch):
    import requests
    http=Mock();http.__enter__=Mock(return_value=http);http.__exit__=Mock(return_value=False)
    info=Mock(status_code=200);info.json.return_value={'capabilities':['completion']}
    reply=Mock(status_code=200);reply.__enter__=Mock(return_value=reply);reply.__exit__=Mock(return_value=False)
    reply.iter_content.return_value=[json.dumps({'message':{'content':json.dumps({'produtos':[{'nome':'Caderno','estoque':5}]})}}).encode()]
    http.post.side_effect=[info,reply]
    monkeypatch.setattr(docs.requests,'Session',lambda:http)
    assert docs._extrair_ollama('produto',[],'modelo:3b','lista')[0]['nome']=='Caderno'
    payload=http.post.call_args.kwargs['json']
    assert payload['format']['type']=='object' and 'tools' not in payload and not payload['stream']
    assert http.trust_env is False
    monkeypatch.setattr(config,'OLLAMA_URL','https://fornecedor.example')
    with pytest.raises(ValueError,match='localhost'):
        docs._extrair_ollama('segredo',[],'modelo:3b','lista')


@pytest.mark.parametrize('nome,conteudo',[('a.xls',b'abc'),('a.xlsx',b'abc'),('a.csv',b'\xff'),('a.pdf',b'abc'),('a.exe',b'abc')])
def test_formatos_invalidos_nao_importam(loja,nome,conteudo):
    with pytest.raises(ValueError):
        docs.ler(arquivo(conteudo,nome),'dono')
    assert estoque.listar_produtos()==[]


def test_pdf_cliente_sem_custo_dados_privados_com_foto_e_paginas(loja):
    estoque.salvar_produto({'nome':'Produto secreto inativo','categoria':'Interno','preco':10,'ativo':False,'estoque':2})
    estoque.salvar_produto({'nome':'Produto esgotado','categoria':'Interno','preco':10,'estoque':0})
    for i in range(25):
        estoque.salvar_produto({'nome':'Caderno pontilhado '+str(i),'categoria':'Papelaria','preco':15.9,'custo':3.14,'sku':str(i),'estoque':7,'descricao':'Descrição longa com acentos. '*8})
    p=estoque.listar_produtos()[0]
    from src.fotos import normalizar
    img=BytesIO();Image.new('RGB',(50,100),'blue').save(img,format='PNG');img.seek(0)
    estoque.salvar_produto({**p,'ativo':True,'controlar_estoque':True},p['id'],_foto=normalizar(img))
    pdf=catalogo_pdf.gerar()
    text=texto_pdf(pdf)
    assert 'Caderno pontilhado 24' in text and '15,90' in text
    assert 'Produto secreto' not in text and 'Produto esgotado' not in text and '3,14' not in text
    assert 'Estoque: 7' not in text
    assert 'Estoque: 7' in texto_pdf(catalogo_pdf.gerar(True))
    with pdfium.PdfDocument(pdf) as d:
        assert len(d)>1


def test_apis_protegidas_e_pdf_publico_com_link_valido(loja):
    app,cliente,headers=loja
    outro=app.test_client()
    assert outro.get('/gestao/api/catalogo.pdf').status_code==401
    assert cliente.post('/gestao/api/documentos/ler',data={}).status_code==403
    r=cliente.post('/gestao/api/documentos/ler',headers=headers,data={'arquivo':(BytesIO(arquivo().stream.read()),'estoque.csv')})
    assert r.status_code==200
    r=cliente.post('/gestao/api/documentos/lancar',headers=headers,json={'token':r.json['token'],'produtos':r.json['produtos']})
    assert r.status_code==200
    from src.links import gerar_link_pedido
    token=gerar_link_pedido('5511999990000').rsplit('/',1)[-1]
    from src.publicacao import PaginaPublica
    publico=Client(PaginaPublica(app,lambda:'https://loja.example'),Response)
    pdf=publico.get('/m/'+token+'/catalogo.pdf')
    assert pdf.status_code==200 and pdf.mimetype=='application/pdf'
    assert publico.get('/gestao/api/catalogo.pdf').status_code==404
    assert publico.post('/gestao/api/documentos/ler').status_code==404
    assert publico.get('/m/invalido/catalogo.pdf').status_code==410
    with get_conn() as c:
        c.execute('UPDATE links_pedido SET usado=1 WHERE token=?',(token,))
    assert publico.get('/m/'+token+'/catalogo.pdf').status_code==410


def test_nota_pendente_pode_ser_salva_sem_repetir_estoque(loja,monkeypatch):
    real=cerebro.registrar_documento_estoque
    monkeypatch.setattr(cerebro,'registrar_documento_estoque',Mock(side_effect=OSError('teste')))
    r=docs.ler(arquivo(),'dono')
    assert lancar(r)['nota_salva'] is False
    monkeypatch.setattr(cerebro,'registrar_documento_estoque',real)
    assert docs.salvar_nota(r['token'],'dono')['ok']
    assert estoque.listar_produtos()[0]['estoque']==5


def test_markdown_obsidian_preenchido_sem_ollama_e_limites(loja,monkeypatch):
    monkeypatch.setattr(docs,'_extrair_ollama',Mock(side_effect=AssertionError('Não usar IA nesta tabela')))
    md='''---
atendimento: false
---
# Lista preenchida
| nome | categoria | preco | sku | estoque |
| --- | --- | --- | --- | --- |
| Camiseta | Roupas | 49,90 | CAM | 2 |
'''
    r=docs.ler(arquivo(md,'folha.md'),'dono')
    assert not r['usou_ollama']
    lancar(r)
    assert estoque.listar_produtos()[0]['sku']=='CAM'
    with pytest.raises(ValueError,match='10 MB'):
        docs.ler(arquivo(b'x'*(docs.LIMITE+1),'folha.pdf'),'dono')


def test_pdf_muito_longo_e_planilha_com_varias_abas_rejeitados(loja):
    out=BytesIO();c=canvas.Canvas(out)
    for i in range(11):
        c.drawString(30,700,'Produto');c.showPage()
    c.save()
    with pytest.raises(ValueError,match='10 páginas'):
        docs.ler(arquivo(out.getvalue(),'nota.pdf'),'dono')
    book=openpyxl.Workbook();book.create_sheet('Segunda');out=BytesIO();book.save(out);book.close()
    with pytest.raises(ValueError,match='única aba'):
        docs.ler(arquivo(out.getvalue(),'nota.xlsx'),'dono')


def test_scans_exigem_visao_e_modelos_cloud_rejeitados(loja,monkeypatch):
    http=Mock();http.__enter__=Mock(return_value=http);http.__exit__=Mock(return_value=False)
    info=Mock(status_code=200);info.json.return_value={'capabilities':['completion']}
    http.post.return_value=info;monkeypatch.setattr(docs.requests,'Session',lambda:http)
    with pytest.raises(ValueError,match='visão'):
        docs._extrair_ollama('', ['base64'], 'modelo:3b','lista')
    with pytest.raises(ValueError,match='modelo local'):
        docs._extrair_ollama('texto', [], 'modelo-cloud','lista')
    info.json.return_value={'remote_host':'https://cloud.example'}
    with pytest.raises(ValueError,match='modelos em nuvem'):
        docs._extrair_ollama('texto', [], 'modelo:3b','lista')


def test_quantidade_inteira_decimal_e_entrada_fracionada(loja):
    r=docs.ler(arquivo('nome;categoria;preco;sku;unidade;estoque\nLapis;Papelaria;2;LAP;un;3.0'),'dono')
    assert r['produtos'][0]['estoque']=='3'
    lancar(r)
    estoque.salvar_produto({'nome':'Café','categoria':'Mercearia','preco':20,'sku':'CAFE','unidade':'kg','estoque':1.5,'fracionado':True})
    r=docs.ler(arquivo('nome;sku;unidade;estoque;fracionado\nCafé;CAFE;kg;0,250;sim'),'dono',tipo='compra')
    lancar(r)
    assert next(p for p in estoque.listar_produtos() if p['sku']=='CAFE')['estoque']==1.75


def test_preco_generico_em_compra_exige_revisar_venda(loja):
    r=docs.ler(arquivo(),'dono',tipo='compra')
    assert r['produtos'][0]['preco']=='' and r['produtos'][0]['custo']=='15,90'
    with pytest.raises(ValueError,match='Preço'):
        lancar(r)
    assert estoque.listar_produtos()==[]
    r['produtos'][0]['preco']='25'
    lancar(r)
    assert estoque.listar_produtos()[0]['preco']==25
